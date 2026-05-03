import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

class WordExporter:
    def __init__(self, font_name="Times New Roman", font_size=12, margin_inches=0.5):
        self.font_name = font_name
        self.font_size = font_size
        self.margin_inches = margin_inches

    def sanitize_text(self, text):
        text = text.replace('\n', ' ').strip()
        text = re.sub(r'\s+', ' ', text)
        return text

    def is_heading(self, line):
        return line.startswith("CHƯƠNG")

    def is_list_item(self, line):
        return bool(re.match(r'^(\-|\+|\d+\.\d+)', line))

    def export_to_docx(self, text_lines, output_path):
        document = Document()

        for section in document.sections:
            section.top_margin = Inches(self.margin_inches)
            section.bottom_margin = Inches(self.margin_inches)
            section.left_margin = Inches(self.margin_inches)
            section.right_margin = Inches(self.margin_inches)

        document_style = document.styles['Normal']
        font = document_style.font
        font.name = self.font_name
        font.size = Pt(self.font_size)

        current_paragraph_chunks = []
        sanitized_lines = [self.sanitize_text(raw) for raw in text_lines if self.sanitize_text(raw)]

        for i, line in enumerate(sanitized_lines):
            is_curr_heading = self.is_heading(line)
            is_curr_list = self.is_list_item(line)

            if is_curr_heading or is_curr_list:
                if current_paragraph_chunks:
                    paragraph = document.add_paragraph(" ".join(current_paragraph_chunks))
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    current_paragraph_chunks = []

                paragraph = document.add_paragraph(line)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if is_curr_heading:
                    for run in paragraph.runs:
                        run.bold = True
            else:
                current_paragraph_chunks.append(line)
                
                # Check next line to apply Level 1 Smart Text Heuristic
                next_line = sanitized_lines[i + 1] if i + 1 < len(sanitized_lines) else None
                
                if next_line is None:
                    # End of lines -> Flush
                    paragraph = document.add_paragraph(" ".join(current_paragraph_chunks))
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    current_paragraph_chunks = []
                else:
                    next_is_heading = self.is_heading(next_line)
                    next_is_list = self.is_list_item(next_line)
                    next_starts_lower = next_line[0].islower()
                    ends_with_punct = line.endswith(('.', ':', '?', '!'))

                    # RULE: Only flush if next line is a structural block OR 
                    # (current ends with punctuation AND next line DOES NOT start with lowercase)
                    if next_is_heading or next_is_list:
                        flush = True
                    elif ends_with_punct and not next_starts_lower:
                        flush = True
                    else:
                        flush = False

                    if flush:
                        paragraph = document.add_paragraph(" ".join(current_paragraph_chunks))
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        current_paragraph_chunks = []

        document.save(output_path)
        return output_path
