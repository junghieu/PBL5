from docx import Document
from docx.shared import Pt
import os


def create_docx_from_text(text_content: str, output_path: str) -> str:
    """
    Create a Word document from extracted text.
    
    Args:
        text_content: The text content to write to the document
        output_path: Path where the document should be saved
        
    Returns:
        Path to the created document
    """
    # Create a new Document
    doc = Document()
    
    # Add a title
    doc.add_heading('OCR Extracted Text', 0)
    
    # Add the extracted text
    # Split by lines and add each as a paragraph
    lines = text_content.split('\n')
    for line in lines:
        if line.strip():  # Only add non-empty lines
            paragraph = doc.add_paragraph(line)
            # Set font size
            for run in paragraph.runs:
                run.font.size = Pt(12)
    
    # Ensure the output directory exists
    output_dir = os.path.dirname(output_path)
    if output_dir:  # Only create if there's a directory in the path
        os.makedirs(output_dir, exist_ok=True)
    
    # Save the document
    doc.save(output_path)
    
    return output_path
